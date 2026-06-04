import os
import time
import warnings
import shutil
import threading
import queue
import sys
import subprocess
import tkinter as tk
from tkinter import font as tkfont
from tkinter import ttk, messagebox
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from openpyxl import load_workbook

from config import PASTA_MONITORADA, PLANILHA_ANALISE, Cores
from utils import imprimir_banner_inicial, registrar_log, realizar_backup_planilha
from excel_manager import mapear_ine_para_abas
from sisab_strategies import ContextoSisab
from pec_strategies import ContextoPec

warnings.filterwarnings("ignore", category=UserWarning)
PASTA_AUTOMACAO = PASTA_MONITORADA
PASTA_PROCESSADOS = os.path.join(PASTA_AUTOMACAO, "Processados")

os.makedirs(PASTA_AUTOMACAO, exist_ok=True)
os.makedirs(PASTA_PROCESSADOS, exist_ok=True)

fila_arquivos = []
fila_interface = queue.Queue()

# Controle de segurança da interface:
# a planilha só é liberada para abertura depois que o lote termina,
# os dados são salvos e os arquivos são movidos para Processados.
PROCESSAMENTO_EM_ANDAMENTO = False
PLANILHA_LIBERADA_PARA_ABRIR = False

def planilha_esta_aberta(caminho_arquivo):
    if not os.path.exists(caminho_arquivo): return False
    try:
        with open(caminho_arquivo, 'a'): pass
        return False
    except PermissionError: return True

def inicializar_mapa_postos():
    if not os.path.exists(PLANILHA_ANALISE): return {}
    try:
        wb = load_workbook(PLANILHA_ANALISE, data_only=True)
        mapa = mapear_ine_para_abas(wb)
        wb.close()
        return mapa
    except Exception as e: return {}

MAPA_NOMES_POSTOS = inicializar_mapa_postos()

class RedirecionadorSaida:
    def __init__(self, fila):
        self.fila = fila
        self.buffer = ""
    def write(self, texto):
        self.buffer += texto
        if '\n' in self.buffer:
            linhas = self.buffer.split('\n')
            for linha in linhas[:-1]:
                if linha.strip(): self.fila.put(linha.strip())
            self.buffer = linhas[-1]
    def flush(self): pass

def limpar_codigos_cor(texto):
    import re
    return re.sub(r'\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@~])', '', texto)

def eh_arquivo_valido(nome_arquivo):
    nome = nome_arquivo.lower()
    nome_planilha_principal = os.path.basename(PLANILHA_ANALISE).lower()
    return not (nome.startswith("~") or nome.startswith(".") or nome == nome_planilha_principal) and nome.endswith((".xlsx", ".pdf", ".csv"))

def processar_fila_em_lote():
    global fila_arquivos, MAPA_NOMES_POSTOS, PROCESSAMENTO_EM_ANDAMENTO, PLANILHA_LIBERADA_PARA_ABRIR
    if not fila_arquivos: return

    if planilha_esta_aberta(PLANILHA_ANALISE):
        print("❌ PAUSA: A planilha 'analise.xlsx' está ABERTA! Feche-a para continuar.")
        return

    lote = list(set(fila_arquivos.copy()))
    fila_arquivos.clear()

    PROCESSAMENTO_EM_ANDAMENTO = True
    PLANILHA_LIBERADA_PARA_ABRIR = False

    print("UI_CMD|LIMPAR_TABELA")
    print(f"⚙️ Processando {len(lote)} arquivo(s)...")

    try:
        realizar_backup_planilha(PLANILHA_ANALISE)

        wb = load_workbook(PLANILHA_ANALISE)
        mapa_abas = mapear_ine_para_abas(wb)

        for caminho in lote:
            if not os.path.exists(caminho): continue
            nome_arquivo_min = os.path.basename(caminho).lower()

            if nome_arquivo_min.endswith(".pdf"):
                ContextoPec(caminho, wb, mapa_abas).executar()
            elif nome_arquivo_min.endswith(".xlsx") and "sisab" in nome_arquivo_min:
                ContextoSisab(
                    caminho,
                    wb,
                    mapa_abas
                ).executar()

        print(f"💾 Salvando resultados na planilha...")
        wb.save(PLANILHA_ANALISE)
        wb.close()

        MAPA_NOMES_POSTOS = inicializar_mapa_postos()
        for caminho in lote:
            if os.path.exists(caminho):
                shutil.move(caminho, os.path.join(PASTA_PROCESSADOS, os.path.basename(caminho)))

        PLANILHA_LIBERADA_PARA_ABRIR = True
        print("✅ Tudo pronto! Dados salvos e arquivos movidos. A planilha já pode ser aberta pelo botão.")

    except Exception as e:
        PLANILHA_LIBERADA_PARA_ABRIR = False
        print(f"❌ Erro crítico: {e}")
        fila_arquivos.extend(lote)

    finally:
        PROCESSAMENTO_EM_ANDAMENTO = False

class MonitorRelatorios(FileSystemEventHandler):
    def on_created(self, event):
        if not event.is_directory and eh_arquivo_valido(os.path.basename(event.src_path)):
            time.sleep(1)
            global PLANILHA_LIBERADA_PARA_ABRIR
            PLANILHA_LIBERADA_PARA_ABRIR = False
            if event.src_path not in fila_arquivos: fila_arquivos.append(event.src_path)

def iniciar_automacao_background():
    global PLANILHA_LIBERADA_PARA_ABRIR
    if os.path.exists(PASTA_AUTOMACAO):
        for arquivo in os.listdir(PASTA_AUTOMACAO):
            caminho_completo = os.path.join(PASTA_AUTOMACAO, arquivo)
            if os.path.isfile(caminho_completo) and eh_arquivo_valido(arquivo) and caminho_completo not in fila_arquivos:
                PLANILHA_LIBERADA_PARA_ABRIR = False
                fila_arquivos.append(caminho_completo)

    observer = Observer()
    observer.schedule(MonitorRelatorios(), path=PASTA_AUTOMACAO, recursive=False)
    observer.start()
    try:
        while True:
            processar_fila_em_lote()
            time.sleep(5)
    except Exception: observer.stop()
    observer.join()

class AplicativoAutomacao:
    def __init__(self, root):
        self.root = root
        self.root.title("Assistente de Análise - Profissional Rosileia")
        self.root.geometry("1100x750")
        self.root.configure(bg="#F4F6F9")

        frame_topo = tk.Frame(self.root, bg="#FFFFFF", pady=15, padx=20, relief="groove", bd=1)
        frame_topo.pack(fill="x", padx=20, pady=15)
        tk.Label(frame_topo, text="👋 Olá, Rosileia!", bg="#FFFFFF", fg="#2C3E50", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(frame_topo, text="O sistema monitora a pasta e salva na planilha automaticamente.", bg="#FFFFFF", fg="#7F8C8D", font=("Segoe UI", 11)).pack(anchor="w")

        caixa_instrucoes = tk.Text(frame_topo, bg="#FFFFFF", font=("Segoe UI", 11), bd=0, highlightthickness=0, height=10, wrap="word")
        caixa_instrucoes.pack(fill="x", pady=5)

        caixa_instrucoes.tag_configure("padrao", foreground="#7F8C8D")
        caixa_instrucoes.tag_configure("alerta", foreground="#E74C3C", font=("Segoe UI", 11, "bold"))
        caixa_instrucoes.tag_configure("ok", foreground="#27AE60", font=("Segoe UI", 11, "bold"))
        caixa_instrucoes.tag_configure("atencao", foreground="#E67E22", font=("Segoe UI", 11, "bold"))

        caixa_instrucoes.insert("end", "1) Para arquivos que vem do PEC o PDF ", "padrao")
        caixa_instrucoes.insert("end", "TEM QUE CONTER O NOME EM MAIÚSCULO", "alerta")
        caixa_instrucoes.insert("end", " da unidade com os acentos.\n", "padrao")
        caixa_instrucoes.insert("end", "2) Para arquivos que vem do SISAB ", "padrao")
        caixa_instrucoes.insert("end", "precisa ser renomeado com Sisab ou sisab", "ok")
        caixa_instrucoes.insert("end", " o nome do arquivo.\n", "padrao")
        caixa_instrucoes.insert("end", "3) Lembre-se de quando baixar os arquivos ", "padrao")
        caixa_instrucoes.insert("end", "DEIXE A PLANILHA FECHADA", "alerta")
        caixa_instrucoes.insert("end", ", abra quando a sua assistente permitir!\n", "padrao")
        caixa_instrucoes.insert("end", "4) Para arquivos do PEC, ", "padrao")
        caixa_instrucoes.insert("end", "APENAS PDF", "atencao")
        caixa_instrucoes.insert("end", " e do SISAB ", "padrao")
        caixa_instrucoes.insert("end", "APENAS arquivo XLSX", "atencao")
        caixa_instrucoes.insert("end", " do EXCEL.\n", "padrao")
        caixa_instrucoes.insert("end", "5) A sua assistente segue o modelo para extrair os dados como consta no ", "padrao")
        caixa_instrucoes.insert("end", "MANUAL", "atencao")
        caixa_instrucoes.insert("end", ", fora disso ele não vai pegar os dados.\n", "padrao")
        caixa_instrucoes.insert("end", "6) Se a planilha possuir algum problema feche a sua assistente, entre na pasta ", "padrao")
        caixa_instrucoes.insert("end", "Cópias", "ok")
        caixa_instrucoes.insert("end", " e pegue a PENULTIMA cópia da planilha na pasta e coloque no lugar da antiga renomeando para 'Análise de Produção' e o ano que você queira ao lado , apos isso abra a sua assistente novamente.\n", "padrao")
        caixa_instrucoes.insert("end", "7) Recomendo usar a assistente APENAS final do ", "padrao")
        caixa_instrucoes.insert("end", "mês ", "ok")
        caixa_instrucoes.insert("end", "para evitar a entrada de dados repetidos! .\n", "padrao")

        caixa_instrucoes.configure(state="disabled")

        frame_status = tk.Frame(self.root, bg="#FFFFFF", pady=10, padx=20, relief="groove", bd=1)
        frame_status.pack(fill="x", padx=20, pady=(0, 15))
        self.var_ultima_acao = tk.StringVar(value="Monitorando a pasta...")
        tk.Label(frame_status, textvariable=self.var_ultima_acao, bg="#FFFFFF", fg="#34495E", font=("Segoe UI", 11)).pack(side="left")

        self.detalhes_por_procedimento = {}

        btn_detalhes = tk.Button(
            frame_status,
            text="🔍 Ver Itens Extraídos",
            bg="#27AE60",
            fg="white",
            font=("Segoe UI", 10, "bold"),
            command=self.abrir_janela_detalhes,
            cursor="hand2",
            relief="flat",
            padx=10
        )
        btn_detalhes.pack(side="right", padx=10)

        btn_abrir_planilha = tk.Button(
            frame_status,
            text="📂 Abrir Planilha",
            bg="#2980B9",
            fg="white",
            font=("Segoe UI", 10, "bold"),
            command=self.abrir_planilha,
            cursor="hand2",
            relief="flat",
            padx=10
        )
        btn_abrir_planilha.pack(side="right", padx=10)

        frame_tabela = tk.Frame(self.root, bg="#F4F6F9")
        frame_tabela.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        colunas = ("ine", "posto", "mes", "indicador", "v_doc", "v_plan", "status")
        self.tabela = ttk.Treeview(frame_tabela, columns=colunas, show="headings")

        self.tabela.heading("ine", text="INE")
        self.tabela.heading("posto", text="Posto (Aba)")
        self.tabela.heading("mes", text="Mês")
        self.tabela.heading("indicador", text="Indicador / Procedimento")
        self.tabela.heading("v_doc", text="Valor Doc.")
        self.tabela.heading("v_plan", text="Valor Planilha")
        self.tabela.heading("status", text="Status na Planilha")

        self.tabela.column("ine", width=80, anchor="center", stretch=False)
        self.tabela.column("posto", width=180, anchor="w", stretch=False)
        self.tabela.column("mes", width=90, anchor="center", stretch=False)
        self.tabela.column("indicador", width=300, anchor="w", stretch=True)
        self.tabela.column("v_doc", width=100, anchor="center", stretch=False)
        self.tabela.column("v_plan", width=120, anchor="center", stretch=False)
        self.tabela.column("status", width=180, anchor="center", stretch=False)

        self.tabela.tag_configure('sucesso', background='#E8F8F5', foreground='#117864')
        self.tabela.tag_configure('erro', background='#FDEDEC', foreground='#C0392B')

        scrollbar_y = ttk.Scrollbar(frame_tabela, orient="vertical", command=self.tabela.yview)
        scrollbar_x = ttk.Scrollbar(frame_tabela, orient="horizontal", command=self.tabela.xview)
        self.tabela.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        scrollbar_y.pack(side="right", fill="y")
        scrollbar_x.pack(side="bottom", fill="x")
        self.tabela.pack(side="left", fill="both", expand=True)

        sys.stdout = RedirecionadorSaida(fila_interface)
        self.atualizar_interface()
        threading.Thread(target=iniciar_automacao_background, daemon=True).start()

    def atualizar_interface(self):
        try:
            while True:
                texto = fila_interface.get_nowait()
                texto_limpo = limpar_codigos_cor(texto)

                if texto_limpo == "UI_CMD|LIMPAR_TABELA":
                    self.detalhes_por_procedimento.clear()
                    for item in self.tabela.get_children():
                        self.tabela.delete(item)

                elif texto_limpo.startswith("UI_DETALHE|"):
                    partes = texto_limpo.split("|")
                    if len(partes) >= 4:
                        proc = partes[1].strip()
                        item = partes[2].strip()
                        val = partes[3].strip()

                        if proc not in self.detalhes_por_procedimento:
                            self.detalhes_por_procedimento[proc] = []

                        linha_formatada = f"{item} - {val}"
                        if linha_formatada not in self.detalhes_por_procedimento[proc]:
                            self.detalhes_por_procedimento[proc].append(linha_formatada)

                elif texto_limpo.startswith("UI_RESULTADO|"):
                    try:
                        partes = texto_limpo.split("|")
                        if len(partes) >= 8:
                            status_cod = partes[1].strip()
                            ine_visual = partes[2].strip()
                            aba_nome = partes[3].strip()
                            indicador = partes[4].strip()
                            v_doc = partes[5].strip()
                            v_plan = partes[6].strip()
                            mes_visual = partes[7].strip()

                            status_map = {
                                "SUCESSO": ("✅ Salvo com sucesso", "sucesso"),
                                "FALHA_MES": ("❌ Mês não encontrado na planilha", "erro"),
                                "FALHA_SISTEMA": ("⚠️ Coluna SISAB/PEC ñ achada", "erro"),
                                "FALHA_LINHA": ("❌ Indicador não achado", "erro"),
                                "FALHA_INE_NAO_MAPEADO": ("❌ INE não cadastrado no sistema", "erro"),
                                "FALHA_ABA_NAO_ENCONTRADA": ("❌ Aba da unidade não existe", "erro"),
                                "FALHA_CATEGORIA": ("⚠️ Categoria não informada", "erro")
                            }
                            texto_status, tag_cor = status_map.get(status_cod, ("❌ Erro desconhecido", "erro"))
                            posto_nome = MAPA_NOMES_POSTOS.get(ine_visual, aba_nome)

                            if ine_visual in ["None", "PDF_SEM_NOME"]: ine_visual = "-"

                            self.tabela.insert("", 0, values=(ine_visual, posto_nome, mes_visual, indicador, v_doc, v_plan, texto_status), tags=(tag_cor,))
                    except Exception as e:
                        pass

                elif not texto_limpo.startswith("DADO_EXTRAIDO|") and not texto_limpo.startswith("UI_"):
                    self.var_ultima_acao.set(texto_limpo)

        except queue.Empty: pass
        self.root.after(200, self.atualizar_interface)

    def abrir_planilha(self):
        if PROCESSAMENTO_EM_ANDAMENTO or fila_arquivos:
            messagebox.showerror(
                "Aguarde a extração",
                "A planilha ainda não pode ser aberta. Aguarde o sistema terminar a extração e salvar os dados."
            )
            return

        if not PLANILHA_LIBERADA_PARA_ABRIR:
            messagebox.showerror(
                "Planilha ainda bloqueada",
                "Aguarde o processo de extração finalizar. Quando aparecer que tudo foi salvo, a planilha poderá ser aberta."
            )
            return

        if not os.path.exists(PLANILHA_ANALISE):
            messagebox.showerror(
                "Planilha não encontrada",
                f"Não encontrei a planilha:\n{PLANILHA_ANALISE}"
            )
            return

        try:
            if sys.platform.startswith("win"):
                os.startfile(PLANILHA_ANALISE)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", PLANILHA_ANALISE])
            else:
                subprocess.Popen(["xdg-open", PLANILHA_ANALISE])
        except Exception as e:
            messagebox.showerror(
                "Erro ao abrir planilha",
                f"Não consegui abrir a planilha automaticamente.\nErro: {e}"
            )

    def abrir_janela_detalhes(self):
        janela = tk.Toplevel(self.root)
        janela.title("Itens Extraídos por Procedimento")
        janela.geometry("760x560")
        janela.configure(bg="#FFFFFF")

        frame_texto = tk.Frame(janela, bg="#FFFFFF")
        frame_texto.pack(fill="both", expand=True, padx=10, pady=10)

        scrollbar_y = ttk.Scrollbar(frame_texto, orient="vertical")
        scrollbar_y.pack(side="right", fill="y")

        txt_area = tk.Text(
            frame_texto,
            font=("Segoe UI", 11),
            bg="#FFFFFF",
            fg="#2C3E50",
            padx=20,
            pady=20,
            relief="flat",
            wrap="word",
            yscrollcommand=scrollbar_y.set
        )
        txt_area.pack(side="left", fill="both", expand=True)
        scrollbar_y.config(command=txt_area.yview)

        if not self.detalhes_por_procedimento:
            txt_area.insert(
                "end",
                "Nenhum dado detalhado foi processado ainda.\n\n"
                "Aguarde a leitura de um arquivo PEC/SISAB."
            )
        else:
            for proc, itens in self.detalhes_por_procedimento.items():
                txt_area.insert("end", f"📌 {proc}\n", "titulo")
                txt_area.insert("end", "\nItens extraídos:\n", "subtitulo")

                for i, item in enumerate(itens, 1):
                    nome_item = item
                    valor_item = ""

                    if " - " in item:
                        nome_item, valor_item = item.rsplit(" - ", 1)

                    txt_area.insert("end", f"\n{i})\n", "numero")
                    txt_area.insert("end", f"   Procedimento/Código: {nome_item}\n")
                    if valor_item:
                        txt_area.insert("end", f"   Valor extraído: {valor_item}\n")

                txt_area.insert("end", "\n" + "─" * 70 + "\n\n")

        txt_area.tag_configure("titulo", font=("Segoe UI", 13, "bold"), spacing1=8, spacing3=4)
        txt_area.tag_configure("subtitulo", font=("Segoe UI", 11, "bold"))
        txt_area.tag_configure("numero", font=("Segoe UI", 11, "bold"))
        txt_area.config(state="disabled")

if __name__ == "__main__":
    root = tk.Tk()
    app = AplicativoAutomacao(root)
    root.mainloop()