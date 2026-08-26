import os
import shutil
from datetime import datetime
from openpyxl import load_workbook
from difflib import SequenceMatcher
from config import PLANILHA_2026, PASTA_COPIAS, MAX_BACKUPS
from utils import padronizar_nome, _tokens_nome

def planilha_esta_aberta(caminho_planilha=PLANILHA_2026):
    if not os.path.exists(caminho_planilha): return False, "Planilha ainda não encontrada."
    pasta = os.path.dirname(caminho_planilha)
    nome = os.path.basename(caminho_planilha)
    lock_excel = os.path.join(pasta, f"~${nome}")
    lock_libreoffice = os.path.join(pasta, f".~lock.{nome}#")
    lock_presente = os.path.exists(lock_excel) or os.path.exists(lock_libreoffice)
    try:
        with open(caminho_planilha, "r+b"): pass
    except PermissionError: return True, "A planilha está aberta. Feche a planilha e lance o PDF novamente."
    except OSError as e: return True, f"Não foi possível acessar a planilha com segurança ({e}). Tente novamente em instantes."

    if lock_presente:
        try:
            with open(os.path.join(pasta, "_aviso_lock_planilha.log"), "a", encoding="utf-8") as f:
                f.write(f"{datetime.now().isoformat()} - Arquivo de lock encontrado mas a planilha abriu normalmente para escrita.\n")
        except OSError: pass
    return False, "Planilha disponível para gravação."

def criar_backup_rotativo(app=None):
    if not os.path.exists(PLANILHA_2026):
        raise FileNotFoundError("Planilha original não encontrada para gerar o backup.")

    os.makedirs(PASTA_COPIAS, exist_ok=True)
    nome_base, ext = os.path.splitext(os.path.basename(PLANILHA_2026))

    # Rotaciona os backups existentes
    for i in range(MAX_BACKUPS - 1, 0, -1):
        caminho_atual = os.path.join(PASTA_COPIAS, f"{nome_base}_backup_{i}{ext}")
        caminho_novo = os.path.join(PASTA_COPIAS, f"{nome_base}_backup_{i+1}{ext}")
        if os.path.exists(caminho_atual):
            shutil.move(caminho_atual, caminho_novo)

    # O novo backup assume sempre a posição 1
    caminho_backup_recente = os.path.join(PASTA_COPIAS, f"{nome_base}_backup_1{ext}")
    shutil.copy2(PLANILHA_2026, caminho_backup_recente)

    if app:
        app.inserir_log(f"BACKUP LOCAL: Cópia de segurança criada (backup_1).", "AZUL")

def _linha_tipo(ws, tipo):
    alvos = ["ATENDIMENTO INDIVIDUAL", "ATIVIDADE INDIVIDUAL"] if tipo == "INDIVIDUAL" else ["ATIVIDADE COLETIVA"]
    for row in range(1, ws.max_row + 1):
        valor = padronizar_nome(ws.cell(row=row, column=1).value)
        if any(alvo in valor for alvo in alvos): return row
    return None

def _coluna_mes_esus(ws, mes_atual):
    mes_pad = padronizar_nome(mes_atual)
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(2, ws.max_column + 1):
            if padronizar_nome(ws.cell(row=row, column=col).value) == mes_pad: return col
    return None

def _eh_aba_profissional(nome_aba):
    nome_pad = padronizar_nome(nome_aba)
    prefixos_ignorados = ("SISAB", "RESUMO", "BASE", "DADOS", "CONFIG", "MODELO", "PLANILHA", "PANILHA", "FOLHA")
    return bool(nome_pad) and not nome_pad.startswith(prefixos_ignorados)

def _pontuar_aba_profissional(aba_pad, nome_pad):
    if not aba_pad or not nome_pad: return 0
    if aba_pad == nome_pad: return 1.0
    tokens_aba = _tokens_nome(aba_pad)
    tokens_nome = _tokens_nome(nome_pad)
    if not tokens_aba or not tokens_nome: return 0
    if all(t in tokens_nome for t in tokens_aba): return 0.96 if len(tokens_aba) == 1 else 0.98
    score = SequenceMatcher(None, aba_pad, nome_pad).ratio()
    if score >= 0.92: return score
    return 0

def _aba_do_profissional(wb, nome_profissional):
    nome_pad = padronizar_nome(nome_profissional)
    abas = {padronizar_nome(nome): nome for nome in wb.sheetnames if _eh_aba_profissional(nome)}
    if nome_pad in abas: return abas[nome_pad]

    candidatos = []
    for aba_pad, aba_original in abas.items():
        score = _pontuar_aba_profissional(aba_pad, nome_pad)
        if score > 0: candidatos.append((score, aba_original))
    candidatos.sort(reverse=True, key=lambda x: x[0])

    if not candidatos or (len(candidatos) > 1 and abs(candidatos[0][0] - candidatos[1][0]) < 0.02):
        return None
    return candidatos[0][1]

def obter_profissionais_planilha(tipo=None, mes_atual=None, nomes_profissionais=None):
    dados_planilha = {}
    if not os.path.exists(PLANILHA_2026): return dados_planilha
    try:
        wb = load_workbook(PLANILHA_2026, read_only=True, data_only=True)
        def valor_da_aba(nome_aba):
            ws = wb[nome_aba]
            if not (tipo and mes_atual): return 0
            linha, coluna = _linha_tipo(ws, tipo), _coluna_mes_esus(ws, mes_atual)
            if linha and coluna:
                val = ws.cell(row=linha, column=coluna).value
                return val if isinstance(val, (int, float)) else 0
            return 0

        if nomes_profissionais:
            for nome_pdf in nomes_profissionais:
                nome_aba = _aba_do_profissional(wb, nome_pdf)
                if nome_aba: dados_planilha[padronizar_nome(nome_pdf)] = valor_da_aba(nome_aba)
        else:
            for nome_aba in wb.sheetnames:
                if _eh_aba_profissional(nome_aba): dados_planilha[padronizar_nome(nome_aba)] = valor_da_aba(nome_aba)
        wb.close()
    except Exception: pass
    return dados_planilha

def atualizar_planilha(tipo, mes_atual, profissionais, app):
    try:
        if not os.path.exists(PLANILHA_2026):
            app.inserir_log(f"ERRO: Planilha não encontrada: {PLANILHA_2026}", "VERMELHO")
            return False

        aberta, _ = planilha_esta_aberta()
        if aberta:
            app.inserir_log("ERRO: Planilha aberta. Feche a planilha e lance o PDF novamente.", "VERMELHO")
            return False

        # Backup foi movido para a camada do orchestrator/interface (lote)

        wb = load_workbook(PLANILHA_2026)
        atualizados = 0
        nao_encontrados = []
        sem_layout = []
        mapa_profissionais = {}

        for nome_profissional in profissionais.keys():
            nome_aba = _aba_do_profissional(wb, nome_profissional)
            if not nome_aba: nao_encontrados.append(nome_profissional)
            else: mapa_profissionais[nome_profissional] = nome_aba

        if nao_encontrados:
            app.inserir_log("AVISO: profissionais fora da planilha serão ignorados: " + ", ".join(nao_encontrados), "AMARELO")

        for nome_profissional, valor_novo in list(profissionais.items()):
            nome_aba = mapa_profissionais.get(nome_profissional)
            if not nome_aba: continue
            ws = wb[nome_aba]
            linha_alvo, coluna_alvo = _linha_tipo(ws, tipo), _coluna_mes_esus(ws, mes_atual)

            if not linha_alvo or not coluna_alvo:
                sem_layout.append(f"{nome_profissional} -> aba '{nome_aba}'")
                continue

            celula = ws.cell(row=linha_alvo, column=coluna_alvo)
            valor_atual = celula.value if isinstance(celula.value, (int, float)) else 0
            celula.value = valor_atual + valor_novo
            atualizados += 1
            app.inserir_log(f"Atualizado: {nome_profissional} -> aba '{nome_aba}', {mes_atual}/E-SUS, {tipo}: +{valor_novo}", "VERDE")

        if atualizados == 0:
            wb.close()
            app.inserir_log("ERRO: Nenhum profissional foi associado à planilha.", "VERMELHO")
            return False

        wb.save(PLANILHA_2026)
        wb.close()

        app.inserir_log("SUCESSO! Atualização concluída. A planilha já pode ser aberta.", "VERDE")
        app.inserir_log(f"Total processado: {atualizados} profissional(is).", "BRANCO")
        return True

    except PermissionError:
        app.inserir_log("ERRO: Planilha aberta. Feche a planilha e lance o PDF novamente.", "VERMELHO")
        return False
    except Exception as e:
        app.inserir_log(f"Erro GERAL: {e}", "VERMELHO")
        return False