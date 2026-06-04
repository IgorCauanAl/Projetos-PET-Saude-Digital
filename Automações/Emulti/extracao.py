import os
import re
import time
import shutil
import pdfplumber
from copy import copy
from difflib import SequenceMatcher
from openpyxl import load_workbook
import unicodedata
from datetime import datetime
from config import PLANILHA_2026, PASTA_COPIAS, MAX_BACKUPS

MESES = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"
]

MESES_NUMERO = {
    "01": "JANEIRO", "1": "JANEIRO",
    "02": "FEVEREIRO", "2": "FEVEREIRO",
    "03": "MARÇO", "3": "MARÇO",
    "04": "ABRIL", "4": "ABRIL",
    "05": "MAIO", "5": "MAIO",
    "06": "JUNHO", "6": "JUNHO",
    "07": "JULHO", "7": "JULHO",
    "08": "AGOSTO", "8": "AGOSTO",
    "09": "SETEMBRO", "9": "SETEMBRO",
    "10": "OUTUBRO", "11": "NOVEMBRO", "12": "DEZEMBRO",
}


def padronizar_nome(nome):
    if nome is None:
        return ""
    nome_str = str(nome)
    nome_sem_acentos = ''.join(
        c for c in unicodedata.normalize('NFD', nome_str)
        if unicodedata.category(c) != 'Mn'
    )
    nome_maiusculo = nome_sem_acentos.upper()
    nome_maiusculo = re.sub(r"[^A-Z0-9\s]", " ", nome_maiusculo)
    return ' '.join(nome_maiusculo.split()).strip()


def copiar_estilo(celula_origem, celula_destino):
    if celula_origem.has_style:
        celula_destino.font = copy(celula_origem.font)
        celula_destino.border = copy(celula_origem.border)
        celula_destino.fill = copy(celula_origem.fill)
        celula_destino.number_format = copy(celula_origem.number_format)
        celula_destino.protection = copy(celula_origem.protection)
        celula_destino.alignment = copy(celula_origem.alignment)


def esperar_download_concluir(caminho_arquivo, timeout=60):
    tempo_decorrido = 0
    tamanho_anterior = -1
    while tempo_decorrido < timeout:
        try:
            if not os.path.exists(caminho_arquivo):
                return False
            tamanho_atual = os.path.getsize(caminho_arquivo)
            if tamanho_atual > 0 and tamanho_atual == tamanho_anterior:
                with open(caminho_arquivo, 'rb'):
                    pass
                return True
            tamanho_anterior = tamanho_atual
        except (OSError, PermissionError):
            pass
        time.sleep(1)
        tempo_decorrido += 1
    return False


def limpar_arquivos_antigos(pasta, horas=24, app=None):
    try:
        agora = time.time()
        for arquivo in os.listdir(pasta):
            caminho_completo = os.path.join(pasta, arquivo)
            if os.path.isfile(caminho_completo):
                tempo_arquivo = os.path.getmtime(caminho_completo)
                if (agora - tempo_arquivo) > (horas * 3600):
                    os.remove(caminho_completo)
                    # Remoção silenciosa para não sobrescrever a mensagem final da operação.
                    pass
    except Exception as e:
        if app:
            app.inserir_log(f"Erro ao limpar: {e}", "VERMELHO")


def _texto_pdf(caminho_pdf):
    texto_total = ""
    with pdfplumber.open(caminho_pdf) as pdf:
        for pagina in pdf.pages:
            texto_pagina = pagina.extract_text() or ""
            texto_total += texto_pagina + "\n"
    return texto_total


def titulo_pdf_valido(caminho_pdf):
    """
    Aceita somente relatórios oficiais de Série Histórica usados pela automação.

    Bloqueia qualquer PDF que não contenha no conteúdo um destes títulos:
    - Relatório de Atividade Coletiva - Série Histórica
    - Relatório de Atendimento Individual - Série Histórica
    - Relatório de Atividade Individual - Série Histórica (aceito como compatibilidade)
    """
    try:
        texto_pad = padronizar_nome(_texto_pdf(caminho_pdf))
    except Exception:
        return False, "Não foi possível ler o PDF para validar o título."

    titulos_validos = [
        # No e-SUS, o relatório individual vem como "atendimento individual",
        # não como "atividade individual".
        "RELATORIO DE ATENDIMENTO INDIVIDUAL SERIE HISTORICA",

        # Coletivo vem como "atividade coletiva".
        "RELATORIO DE ATIVIDADE COLETIVA SERIE HISTORICA",

        # Mantido por compatibilidade caso algum relatório/exportação venha com esse texto.
        "RELATORIO DE ATIVIDADE INDIVIDUAL SERIE HISTORICA",
    ]

    if any(titulo in texto_pad for titulo in titulos_validos):
        return True, "PDF válido."

    return False, (
        "PDF rejeitado: o título precisa ser 'Relatório de Atividade Coletiva - Série Histórica' "
        "ou 'Relatório de Atendimento Individual - Série Histórica'."
    )


def planilha_esta_aberta(caminho_planilha=PLANILHA_2026):
    """
    Verifica sinais de que a planilha está aberta no Excel/LibreOffice.
    Retorna (True, mensagem) quando houver risco de falha na gravação.
    """
    if not os.path.exists(caminho_planilha):
        return False, "Planilha ainda não encontrada."

    pasta = os.path.dirname(caminho_planilha)
    nome = os.path.basename(caminho_planilha)
    temporarios = [
        os.path.join(pasta, f"~${nome}"),
        os.path.join(pasta, f".~lock.{nome}#"),
    ]
    for temp in temporarios:
        if os.path.exists(temp):
            return True, "A planilha parece estar aberta. Feche o Excel/LibreOffice e lance o PDF novamente."

    # Teste conservador: no Windows, renomear arquivo aberto pelo Excel costuma falhar.
    try:
        os.rename(caminho_planilha, caminho_planilha)
    except PermissionError:
        return True, "A planilha está aberta. Feche a planilha e lance o PDF novamente."
    except OSError:
        # Em alguns ambientes, os.rename para o mesmo nome pode não ser confiável.
        # Não bloqueia se não for claramente PermissionError.
        pass

    return False, "Planilha disponível para gravação."


def criar_backup_rotativo(app=None):
    """Cria backup com data/hora e mantém somente as 5 cópias mais recentes."""
    if not os.path.exists(PLANILHA_2026):
        return

    os.makedirs(PASTA_COPIAS, exist_ok=True)
    nome_base, ext = os.path.splitext(os.path.basename(PLANILHA_2026))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho_backup = os.path.join(PASTA_COPIAS, f"{nome_base}_backup_{timestamp}{ext}")
    shutil.copy2(PLANILHA_2026, caminho_backup)

    backups = []
    prefixo = f"{nome_base}_backup_"
    for arquivo in os.listdir(PASTA_COPIAS):
        if arquivo.startswith(prefixo) and arquivo.lower().endswith(ext.lower()):
            caminho = os.path.join(PASTA_COPIAS, arquivo)
            backups.append((os.path.getmtime(caminho), caminho))

    backups.sort(reverse=True)
    for _, caminho_antigo in backups[MAX_BACKUPS:]:
        try:
            os.remove(caminho_antigo)
        except OSError:
            pass

    if app:
        app.inserir_log(f"BACKUP LOCAL: cópia criada. Mantendo até {MAX_BACKUPS} backups.", "AZUL")


def _mes_por_numero(mes_num):
    try:
        return MESES_NUMERO.get(str(int(str(mes_num).strip())))
    except Exception:
        return None


def _extrair_mes(texto_total):
    """
    Extrai o mês/competência de forma conservadora.

    Prioriza linhas de competência/período para não confundir com data de
    impressão do relatório. Quando encontrar uma data no padrão dd/mm/aaaa,
    usa o segundo campo como mês.
    """
    linhas = texto_total.splitlines()

    palavras_chave = [
        "PERIODO", "COMPETENCIA", "MES", "MÊS", "DATA INICIAL",
        "DATA DE INICIO", "DATA INICIO"
    ]

    # 1) Primeiro procura em linhas específicas do cabeçalho/competência.
    for linha in linhas:
        linha_pad = padronizar_nome(linha)
        if not any(padronizar_nome(chave) in linha_pad for chave in palavras_chave):
            continue

        for mes in MESES:
            if re.search(rf"\b{mes}\b", linha_pad):
                return mes

        # Ex.: Período: 01/05/2026 a 31/05/2026
        m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", linha)
        if m:
            mes = _mes_por_numero(m.group(2))
            if mes:
                return mes

        # Ex.: Competência: 05/2026
        m = re.search(r"(?<!\d)(\d{1,2})[/-](\d{2,4})(?!\d)", linha)
        if m:
            mes = _mes_por_numero(m.group(1))
            if mes:
                return mes

    # 2) Procura por mês escrito por extenso perto de termos de produção.
    for linha in linhas:
        linha_pad = padronizar_nome(linha)
        if "IMPRESSO" in linha_pad or "PAGINA" in linha_pad or linha_pad.startswith("PAG "):
            continue
        for mes in MESES:
            if re.search(rf"\b{mes}\b", linha_pad):
                return mes

    # 3) Último recurso: primeira data que não esteja em linha de impressão/rodapé.
    for linha in linhas:
        linha_pad = padronizar_nome(linha)
        if "IMPRESSO" in linha_pad or "PAGINA" in linha_pad or linha_pad.startswith("PAG "):
            continue
        m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", linha)
        if m:
            mes = _mes_por_numero(m.group(2))
            if mes:
                return mes

    return "MÊS NÃO IDENTIFICADO"

def _extrair_tipo_relatorio(texto_total):
    texto_pad = padronizar_nome(texto_total)
    if "ATIVIDADE COLETIVA" in texto_pad or "RELATORIO DE ATIVIDADE COLETIVA" in texto_pad:
        return "COLETIVO"
    if "ATENDIMENTO INDIVIDUAL" in texto_pad or "RELATORIO DE ATENDIMENTO INDIVIDUAL" in texto_pad:
        return "INDIVIDUAL"
    return "DESCONHECIDO"


def _limpar_nome_profissional(nome):
    nome = re.sub(r"\s+", " ", str(nome)).strip()

    # Remove valor/quantidade grudado no fim do nome.
    # Ex.: "MARIA DA SILVA 12" -> "MARIA DA SILVA"
    nome = re.sub(r"\s+\d+[\d\.,]*\s*$", "", nome).strip()

    # Remove fragmentos comuns que podem vir grudados no nome durante a extração.
    cortes = [
        "CBO", "CNS", "CPF", "INE", "CNES", "TOTAL", "QUANTIDADE",
        "PROCEDIMENTO", "OCUPAÇÃO", "OCUPACAO", "EQUIPE", "UNIDADE"
    ]
    nome_pad = padronizar_nome(nome)
    for corte in cortes:
        pos = nome_pad.find(corte)
        if pos > 0:
            nome = nome[:pos].strip()
            break

    nome_pad = padronizar_nome(nome)
    # Remove números soltos que tenham sobrado no meio do nome.
    nome_pad = re.sub(r"\b\d+\b", " ", nome_pad)
    return ' '.join(nome_pad.split()).strip()


def _nome_profissional_valido(nome, origem_tabela=True):
    """
    Evita que pedaços quebrados do PDF, como ELA ou ROS, apareçam como
    profissional na interface.
    """
    nome_pad = padronizar_nome(nome)
    if _eh_nome_resumo_ou_total(nome_pad):
        return False
    if len(nome_pad) < 5:
        return False
    if re.search(r"\d", nome_pad):
        return False

    tokens = _tokens_nome(nome_pad) if '_tokens_nome' in globals() else [t for t in nome_pad.split() if t not in {"DE", "DA", "DAS", "DO", "DOS", "E"}]
    if not tokens:
        return False

    palavras_invalidas = {
        "ELA", "ROS", "PROFISSIONAL", "QUANTIDADE", "VALOR", "TOTAL",
        "EQUIPE", "MUNICIPIO", "UNIDADE", "RELATORIO", "PAGINA", "IMPRESSO"
    }
    if nome_pad in palavras_invalidas or any(t in palavras_invalidas for t in tokens):
        return False

    # Em linhas de tabela, o nome completo do profissional geralmente vem com
    # dois ou mais tokens. Isso bloqueia fragmentos como "ELA" e "ROS".
    if origem_tabela and len(tokens) < 2:
        return False

    return True


def _parse_valor(valor_texto):
    valor_limpo = re.sub(r"[^0-9]", "", str(valor_texto))
    if not valor_limpo:
        return None
    return int(valor_limpo)


def _eh_nome_resumo_ou_total(nome):
    """
    Identifica linhas de resumo que não podem ser tratadas como profissional.

    Exemplo do problema: o PDF pode trazer "Total da equipe 123" na mesma
    região da tabela. Essa linha representa o valor total do relatório, não o
    nome de um profissional.
    """
    nome_pad = padronizar_nome(nome)
    if not nome_pad:
        return True

    termos_resumo = [
        "TOTAL DA EQUIPE",
        "TOTAL EQUIPE",
        "TOTAL DE EQUIPE",
        "TOTAL GERAL",
        "TOTAL DO PROFISSIONAL",
        "TOTAL PROFISSIONAL",
        "SUBTOTAL",
        "RESUMO",
    ]
    if any(termo in nome_pad for termo in termos_resumo):
        return True

    # Qualquer linha cujo primeiro token é TOTAL é resumo, não pessoa.
    # Isso evita que "TOTAL DA EQUIPE" apareça na interface como profissional.
    if nome_pad.startswith("TOTAL ") or nome_pad == "TOTAL":
        return True

    return False


def _extrair_total_geral(linhas):
    for linha in linhas:
        linha_pad = padronizar_nome(linha)
        if "TOTAL GERAL" in linha_pad or linha_pad.startswith("TOTAL"):
            m = re.search(r"(\d[\d\.\,]*)\s*$", linha.strip())
            if m:
                return _parse_valor(m.group(1))
    return None


def _extrair_profissional_unico(linhas):
    """Extrai o profissional quando o PDF é um relatório filtrado por um profissional."""
    padroes = [
        r"PROFISSIONAL\s*:?\s*(.+)",
        r"NOME DO PROFISSIONAL\s*:?\s*(.+)",
    ]

    for linha in linhas:
        linha_sem = linha.strip()
        linha_pad = padronizar_nome(linha_sem)
        if not linha_pad or linha_pad == "PROFISSIONAL":
            continue

        for padrao in padroes:
            m = re.search(padrao, linha_sem, flags=re.IGNORECASE)
            if not m:
                continue
            candidato = m.group(1).strip()
            # Evita pegar cabeçalho de tabela: "Profissional Total".
            cand_pad = padronizar_nome(candidato)
            if cand_pad in {"TOTAL", "QUANTIDADE", "PRODUCAO", "VALOR"}:
                continue
            if _eh_nome_resumo_ou_total(cand_pad):
                continue
            nome = _limpar_nome_profissional(candidato)
            if not _nome_profissional_valido(nome, origem_tabela=False):
                continue
            return nome
    return None


def _extrair_profissionais_tabela(linhas):
    """Extrai linhas do tipo: NOME DO PROFISSIONAL 123."""
    profissionais = {}
    nomes_duplicados = set()
    processando = False

    termos_ignorar = [
        "IMPRESSO EM", "PAG", "PAGINA", "RELATORIO", "MUNICIPIO",
        "UNIDADE DE SAUDE", "NIVEL DE DETALHE", "PERIODO", "CNES",
        "PROFISSIONAL TOTAL", "PROFISSIONAL QUANTIDADE"
    ]

    for linha in linhas:
        linha_original = linha.strip()
        linha_pad = padronizar_nome(linha_original)
        if not linha_pad:
            continue
        if "TOTAL GERAL" in linha_pad:
            break
        # Linhas de resumo devem fornecer valor total no fallback, mas nunca
        # podem virar "profissional" na tabela da interface.
        if _eh_nome_resumo_ou_total(linha_pad):
            continue
        if any(t in linha_pad for t in termos_ignorar):
            if "PROFISSIONAL" in linha_pad and ("TOTAL" in linha_pad or "QUANTIDADE" in linha_pad):
                processando = True
            continue
        if linha_pad == "PROFISSIONAL" or linha_pad.startswith("PROFISSIONAL "):
            processando = True
            continue
        if _eh_nome_resumo_ou_total(linha_pad):
            continue

        m = re.match(r"^(.+?)\s+(\d[\d\.\,]*)$", linha_original)
        if not (processando and m):
            continue

        nome = _limpar_nome_profissional(m.group(1))
        valor = _parse_valor(m.group(2))
        if not nome or valor is None:
            continue
        if not _nome_profissional_valido(nome, origem_tabela=True):
            continue

        if nome in profissionais:
            profissionais[nome] += valor
            nomes_duplicados.add(nome)
        else:
            profissionais[nome] = valor

    return profissionais, nomes_duplicados


def ler_relatorio_pdf(caminho_pdf):
    valido, motivo = titulo_pdf_valido(caminho_pdf)
    if not valido:
        raise ValueError(motivo)

    texto_total = _texto_pdf(caminho_pdf)
    texto_upper = texto_total.upper()
    texto_pad = padronizar_nome(texto_total)

    # Validação estrutural: mantém fail-safe, mas menos frágil para variações do relatório.
    marcadores_minimos = ["MINISTERIO DA SAUDE", "SANTO ANTONIO DE JESUS"]
    for marcador in marcadores_minimos:
        if marcador not in texto_pad:
            raise ValueError(f"Estrutura inválida ou arquivo incorreto. Marcador ausente: '{marcador}'")

    if "NIVEL DE DETALHE" in texto_pad and "PROFISSIONAL" not in texto_pad:
        raise ValueError("Estrutura inválida. O relatório não está detalhado por profissional.")

    linhas = texto_total.splitlines()
    mes_atual = _extrair_mes(texto_total)
    tipo = _extrair_tipo_relatorio(texto_total)

    if tipo == "DESCONHECIDO":
        raise ValueError("Estrutura inválida. O PDF não é de Atividade Coletiva nem Individual.")
    if mes_atual == "MÊS NÃO IDENTIFICADO":
        raise ValueError("Não foi possível identificar o mês/competência do PDF.")

    profissionais, nomes_duplicados = _extrair_profissionais_tabela(linhas)

    # Defesa em profundidade: remove qualquer resumo que tenha passado pelo parser.
    profissionais = {
        nome: valor
        for nome, valor in profissionais.items()
        if not _eh_nome_resumo_ou_total(nome)
    }
    nomes_duplicados = {
        nome for nome in nomes_duplicados
        if not _eh_nome_resumo_ou_total(nome)
    }

    # Fallback: relatório filtrado por um profissional, com total geral no final.
    if not profissionais:
        nome_unico = _extrair_profissional_unico(linhas)
        total = _extrair_total_geral(linhas)
        if nome_unico and total is not None:
            profissionais[nome_unico] = total

    if not profissionais:
        raise ValueError("Não foi possível identificar profissional e valor no PDF.")

    return tipo, mes_atual, profissionais, nomes_duplicados


def _linha_tipo(ws, tipo):
    alvo = "ATIVIDADE INDIVIDUAL" if tipo == "INDIVIDUAL" else "ATIVIDADE COLETIVA"
    for row in range(1, ws.max_row + 1):
        valor = padronizar_nome(ws.cell(row=row, column=1).value)
        if valor == alvo or alvo in valor:
            return row
    return None


def _coluna_mes_esus(ws, mes_atual):
    mes_pad = padronizar_nome(mes_atual)
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(1, ws.max_column + 1):
            valor_mes = padronizar_nome(ws.cell(row=row, column=col).value)
            valor_sub = padronizar_nome(ws.cell(row=row + 1, column=col).value)
            if valor_mes == mes_pad and valor_sub == "E SUS":
                return col
            # Fallback caso a linha E-SUS seja removida, mas o mês esteja no cabeçalho.
            if valor_mes == mes_pad and col >= 2:
                return col
    return None


def _mapa_abas_profissionais(wb):
    return {padronizar_nome(nome): nome for nome in wb.sheetnames}


def _eh_aba_profissional(nome_aba):
    """Ignora abas auxiliares que não representam profissionais da gestora."""
    nome_pad = padronizar_nome(nome_aba)
    prefixos_ignorados = ("SISAB", "RESUMO", "BASE", "DADOS", "CONFIG", "MODELO")
    return bool(nome_pad) and not nome_pad.startswith(prefixos_ignorados)


def _tokens_nome(nome):
    conectores = {"DE", "DA", "DAS", "DO", "DOS", "E"}
    return [t for t in padronizar_nome(nome).split() if t not in conectores]


def _pontuar_aba_profissional(aba_pad, nome_pad):
    """
    Pontuação conservadora para evitar associar profissional errado.

    Regras principais:
    - correspondência exata recebe 1.0;
    - aba com primeiro nome/apelido só é aceita quando o token aparece inteiro no nome do PDF;
    - substring solta foi removida para evitar falso positivo, exemplo: ANA dentro de MARIANA;
    - fuzzy match só entra com limite alto.
    """
    if not aba_pad or not nome_pad:
        return 0
    if aba_pad == nome_pad:
        return 1.0

    tokens_aba = _tokens_nome(aba_pad)
    tokens_nome = _tokens_nome(nome_pad)
    if not tokens_aba or not tokens_nome:
        return 0

    # Ex.: aba "Jéssica" e PDF "JESSICA MARIA ...".
    if all(t in tokens_nome for t in tokens_aba):
        return 0.96 if len(tokens_aba) == 1 else 0.98

    # Ex.: aba com nome completo e PDF com nome levemente diferente.
    score = SequenceMatcher(None, aba_pad, nome_pad).ratio()
    if score >= 0.92:
        return score

    return 0


def _aba_do_profissional(wb, nome_profissional):
    nome_pad = padronizar_nome(nome_profissional)
    abas = {
        padronizar_nome(nome): nome
        for nome in wb.sheetnames
        if _eh_aba_profissional(nome)
    }

    if nome_pad in abas:
        return abas[nome_pad]

    candidatos = []
    for aba_pad, aba_original in abas.items():
        score = _pontuar_aba_profissional(aba_pad, nome_pad)
        if score > 0:
            candidatos.append((score, aba_original))

    candidatos.sort(reverse=True, key=lambda x: x[0])
    if not candidatos:
        return None

    # Se duas abas empatam quase igual, não arrisca atualizar a aba errada.
    if len(candidatos) > 1 and abs(candidatos[0][0] - candidatos[1][0]) < 0.02:
        return None

    return candidatos[0][1]


def obter_profissionais_planilha(tipo=None, mes_atual=None, nomes_profissionais=None):
    """
    Retorna o valor atual da célula E-SUS do mês/tipo.

    - Se `nomes_profissionais` for informado, retorna APENAS os nomes que possuem
      aba autorizada na planilha. Nomes ausentes ficam fora do dicionário para
      permitir o bloqueio seguro na interface antes de alterar o Excel.
    - Se não for informado, retorna uma visão geral por aba profissional da planilha.
    """
    dados_planilha = {}
    if not os.path.exists(PLANILHA_2026):
        return dados_planilha

    try:
        wb = load_workbook(PLANILHA_2026, read_only=True, data_only=True)

        def valor_da_aba(nome_aba):
            ws = wb[nome_aba]
            if not (tipo and mes_atual):
                return 0
            linha = _linha_tipo(ws, tipo)
            coluna = _coluna_mes_esus(ws, mes_atual)
            if linha and coluna:
                val = ws.cell(row=linha, column=coluna).value
                return val if isinstance(val, (int, float)) else 0
            return 0

        if nomes_profissionais:
            for nome_pdf in nomes_profissionais:
                nome_aba = _aba_do_profissional(wb, nome_pdf)
                if not nome_aba:
                    continue
                chave = padronizar_nome(nome_pdf)
                dados_planilha[chave] = valor_da_aba(nome_aba)
        else:
            for nome_aba in wb.sheetnames:
                if not _eh_aba_profissional(nome_aba):
                    continue
                dados_planilha[padronizar_nome(nome_aba)] = valor_da_aba(nome_aba)
        wb.close()
    except Exception:
        pass
    return dados_planilha


def atualizar_planilha(tipo, mes_atual, profissionais, app):
    try:
        if not os.path.exists(PLANILHA_2026):
            app.inserir_log(f"ERRO: Planilha não encontrada: {PLANILHA_2026}", "VERMELHO")
            return False

        aberta, mensagem_aberta = planilha_esta_aberta()
        if aberta:
            app.inserir_log("ERRO: Planilha aberta. Feche a planilha e lance o PDF novamente.", "VERMELHO")
            return False

        wb = load_workbook(PLANILHA_2026)
        atualizados = 0
        sem_layout = []

        # Validação por profissional: nomes fora da planilha NÃO bloqueiam mais
        # quem está correto. Eles são ignorados e registrados no log.
        mapa_profissionais = {}
        nao_encontrados = []
        for nome_profissional in profissionais.keys():
            nome_aba = _aba_do_profissional(wb, nome_profissional)
            if not nome_aba:
                nao_encontrados.append(nome_profissional)
            else:
                mapa_profissionais[nome_profissional] = nome_aba

        if nao_encontrados:
            app.inserir_log(
                "AVISO: profissionais fora da planilha serão ignorados: " + ", ".join(nao_encontrados),
                "AMARELO"
            )

        for nome_profissional, valor_novo in list(profissionais.items()):
            nome_aba = mapa_profissionais.get(nome_profissional)
            if not nome_aba:
                # Pula somente o profissional não autorizado, sem impedir os demais.
                continue
            ws = wb[nome_aba]
            linha_alvo = _linha_tipo(ws, tipo)
            coluna_alvo = _coluna_mes_esus(ws, mes_atual)

            if not linha_alvo or not coluna_alvo:
                sem_layout.append(f"{nome_profissional} -> aba '{nome_aba}'")
                continue

            celula_alvo = ws.cell(row=linha_alvo, column=coluna_alvo)
            valor_atual = celula_alvo.value if isinstance(celula_alvo.value, (int, float)) else 0
            celula_alvo.value = valor_atual + valor_novo
            atualizados += 1
            app.inserir_log(
                f"Atualizado: {nome_profissional} -> aba '{nome_aba}', {mes_atual}/E-SUS, {tipo}: +{valor_novo}",
                "VERDE"
            )

        if atualizados == 0:
            wb.close()
            app.inserir_log("ERRO: Nenhum profissional foi associado à planilha.", "VERMELHO")
            if nao_encontrados:
                app.inserir_log("Não encontrados: " + ", ".join(nao_encontrados), "AMARELO")
            if sem_layout:
                app.inserir_log("Abas sem layout esperado: " + ", ".join(sem_layout), "AMARELO")
            return False

        wb.save(PLANILHA_2026)
        wb.close()

        try:
            criar_backup_rotativo(app)
        except Exception as e:
            app.inserir_log(f"Aviso - Falha ao criar backup: {e}", "AMARELO")

        app.inserir_log("SUCESSO! Atualização concluída. A planilha já pode ser aberta.", "VERDE")
        app.inserir_log(f"Total processado: {atualizados} profissional(is).", "BRANCO")

        if nao_encontrados:
            app.inserir_log("Profissionais sem aba correspondente: " + ", ".join(nao_encontrados), "AMARELO")
        if sem_layout:
            app.inserir_log("Abas sem mês/E-SUS ou linha do tipo: " + ", ".join(sem_layout), "AMARELO")
        return True

    except PermissionError:
        app.inserir_log("ERRO: Planilha aberta. Feche a planilha e lance o PDF novamente.", "VERMELHO")
        return False
    except Exception as e:
        app.inserir_log(f"Erro GERAL: {e}", "VERMELHO")
        return False
