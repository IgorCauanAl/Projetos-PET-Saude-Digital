import os
import time
import shutil
import pdfplumber
import pandas as pd
from copy import copy
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from datetime import datetime
import unicodedata

# Importações para a Interface Gráfica
import customtkinter as ctk
import tkinter as tk
from tkinter import messagebox
import queue
import threading

# === CONFIGURAÇÕES DE CAMINHOS ===
PASTA_MONITORADA = r"C:\Users\win10\Desktop\multiteste"
PASTA_DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
PLANILHA_2026 = os.path.join(PASTA_MONITORADA, "Planilha Produção 2026 Teste.xlsx")
PASTA_PROCESSADOS = os.path.join(PASTA_MONITORADA, "processados")

if not os.path.exists(PASTA_PROCESSADOS):
    os.makedirs(PASTA_PROCESSADOS)

# === FUNÇÕES AUXILIARES DE PROCESSAMENTO ===

def padronizar_nome(nome):
    if not nome:
        return ""
    nome_str = str(nome)
    nome_sem_acentos = ''.join(c for c in unicodedata.normalize('NFD', nome_str) if unicodedata.category(c) != 'Mn')
    nome_maiusculo = nome_sem_acentos.upper()
    nome_sem_espacos_duplos = ' '.join(nome_maiusculo.split())
    return nome_sem_espacos_duplos.strip()

def copiar_estilo(celula_origem, celula_destino):
    if celula_origem.has_style:
        celula_destino.font = copy(celula_origem.font)
        celula_destino.border = copy(celula_origem.border)
        celula_destino.fill = copy(celula_origem.fill)
        celula_destino.number_format = copy(celula_origem.number_format)
        celula_destino.protection = copy(celula_origem.protection)
        celula_destino.alignment = copy(celula_origem.alignment)

def esperar_download_concluir(caminho_arquivo, timeout=60):
    # NOVA LÓGICA: Verifica se o tamanho do arquivo estabilizou (parou de crescer)
    tempo_decorrido = 0
    tamanho_anterior = -1
    
    while tempo_decorrido < timeout:
        try:
            if not os.path.exists(caminho_arquivo):
                return False
                
            tamanho_atual = os.path.getsize(caminho_arquivo)
            
            # Se o tamanho é maior que 0 e não mudou desde a última checagem (1 segundo atrás)
            if tamanho_atual > 0 and tamanho_atual == tamanho_anterior:
                # Tenta abrir apenas em modo de leitura para confirmar acesso
                with open(caminho_arquivo, 'rb'):
                    pass
                return True
                
            tamanho_anterior = tamanho_atual
        except (OSError, PermissionError):
            pass # Arquivo ainda está bloqueado pelo navegador
            
        time.sleep(1) 
        tempo_decorrido += 1
        
    return False 

def limpar_arquivos_antigos(pasta, horas=24, app=None):
    try:
        agora = time.time()
        for arquivo in os.listdir(pasta):
            caminho_completo = os.path.join(pasta, arquivo)
            if os.path.isfile(caminho_completo):
                tempo_arquivo = os.path.getmtime(caminho_completo)
                if (agora - tempo_arquivo) > (horas * 3600):
                    os.remove(caminho_completo)
                    if app:
                        app.inserir_log(f"Arquivo expirado removido: {arquivo}", "AMARELO")
    except Exception as e:
        if app:
            app.inserir_log(f"Erro ao tentar limpar arquivos antigos: {e}", "VERMELHO")

def obter_profissionais_planilha():
    existentes = set()
    if not os.path.exists(PLANILHA_2026):
        return existentes
    try:
        wb = load_workbook(PLANILHA_2026, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=13, max_col=1):
            cell_value = row[0].value
            if not cell_value: continue
            nome_padronizado = padronizar_nome(str(cell_value))
            if "TOTAL" in nome_padronizado: break 
            if nome_padronizado:
                existentes.add(nome_padronizado)
        wb.close()
    except Exception:
        pass
    return existentes

def ler_relatorio_pdf(caminho_pdf):
    texto_total = ""
    with pdfplumber.open(caminho_pdf) as pdf:
        for pagina in pdf.pages:
            texto_total += pagina.extract_text() + "\n"
    
    linhas = texto_total.splitlines()
    periodo_linha = next((l for l in linhas if "Período:" in l), "")
    mes_atual = "MÊS NÃO IDENTIFICADO"
    if periodo_linha:
        try:
            data_inicio = periodo_linha.split("Período:")[1].split("a")[0].strip()
            mes_num = pd.to_datetime(data_inicio, dayfirst=True, errors='coerce').month
            meses = ["JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
            if mes_num:
                mes_atual = meses[mes_num - 1]
        except Exception:
            pass

    tipo = "DESCONHECIDO"
    if "Relatório de atividade coletiva" in texto_total:
        tipo = "COLETIVO"
    elif "Relatório de atendimento individual" in texto_total:
        tipo = "INDIVIDUAL"

    profissionais = {}
    nomes_duplicados = set() 
    processando_profissionais = False 
    termos_ignorados = ["IMPRESSO EM", "PAG.", "PAGINA", "RELATORIO"]
    
    for linha in linhas:
        linha_padronizada = padronizar_nome(linha)
        if "TOTAL GERAL" in linha_padronizada: break 
        if "PROFISSIONAL" in linha_padronizada:
            processando_profissionais = True
            continue 
        if "TOTAL" in linha_padronizada: continue
        if any(termo in linha_padronizada for termo in termos_ignorados): continue

        if processando_profissionais:
            if any(c.isdigit() for c in linha.strip().split()[-1:]):
                partes = linha.strip().rsplit(" ", 1)
                if len(partes) == 2:
                    nome, valor_texto = partes
                    try:
                        if nome.strip().endswith(valor_texto):
                            nome = nome.strip()[:-len(valor_texto)].strip()
                        valor_numerico = int(valor_texto)
                        nome_padronizado = padronizar_nome(nome)
                        
                        if nome_padronizado: 
                            if nome_padronizado in profissionais:
                                profissionais[nome_padronizado] += valor_numerico
                                nomes_duplicados.add(nome_padronizado)
                            else:
                                profissionais[nome_padronizado] = valor_numerico
                    except ValueError: continue
    return tipo, mes_atual, profissionais, nomes_duplicados

def atualizar_planilha(tipo, mes_atual, profissionais, app):
    try:
        wb = load_workbook(PLANILHA_2026)
        ws = wb.active

        col_mes_inicio = None
        for cell in ws[11]:
            if cell.value and padronizar_nome(str(cell.value)) == mes_atual:
                col_mes_inicio = cell.column
                break
        
        if not col_mes_inicio:
            app.inserir_log(f"ERRO: Mês '{mes_atual}' não encontrado.", "VERMELHO")
            wb.close()
            return False
        
        coluna_alvo = None
        if tipo == "INDIVIDUAL":
            if padronizar_nome(str(ws.cell(row=12, column=col_mes_inicio).value)) == "INDIVIDUAL":
                coluna_alvo = col_mes_inicio
        elif tipo == "COLETIVO":
            if padronizar_nome(str(ws.cell(row=12, column=col_mes_inicio + 1).value)) == "COLETIVO":
                coluna_alvo = col_mes_inicio + 1
        
        if not coluna_alvo:
            app.inserir_log(f"ERRO: Coluna '{tipo}' não encontrada.", "VERMELHO")
            wb.close()
            return False

        atualizados = 0
        novos_inseridos = [] 
        linha_insercao = None
        
        for row in ws.iter_rows(min_row=13, max_col=1):
            valor = padronizar_nome(str(row[0].value))
            if "TOTAL" in valor: 
                linha_insercao = row[0].row
                break
        
        if not linha_insercao:
            linha_insercao = ws.max_row + 1

        for row in ws.iter_rows(min_row=13, max_row=linha_insercao-1, max_col=1): 
            cell_nome = row[0]
            if not cell_nome.value: continue
            nome_planilha = padronizar_nome(str(cell_nome.value))
            if nome_planilha in profissionais:
                celula_alvo = ws.cell(row=cell_nome.row, column=coluna_alvo)
                valor_atual = celula_alvo.value
                if not isinstance(valor_atual, (int, float)):
                    valor_atual = 0
                celula_alvo.value = valor_atual + profissionais[nome_planilha]
                atualizados += 1
                del profissionais[nome_planilha] 

        if profissionais:
            app.inserir_log("Inserindo novos profissionais na linha de baixo...", "CIANO")
            for nome_novo, valor_novo in profissionais.items():
                ws.insert_rows(linha_insercao)
                ws.cell(row=linha_insercao, column=1, value=nome_novo)
                ws.cell(row=linha_insercao, column=coluna_alvo, value=valor_novo)
                
                linha_referencia = linha_insercao - 1
                for col in range(1, ws.max_column + 1):
                    cell_ref = ws.cell(row=linha_referencia, column=col)
                    cell_nova = ws.cell(row=linha_insercao, column=col)
                    copiar_estilo(cell_ref, cell_nova)
                    if cell_ref.data_type == 'f' and not cell_nova.value:
                        try:
                            formula_original = cell_ref.value
                            cell_nova.value = Translator(formula_original, origin=cell_ref.coordinate).translate_formula(cell_nova.coordinate)
                        except: pass 
                
                novos_inseridos.append(f"{nome_novo} -> {valor_novo}")
                linha_insercao += 1 
                atualizados += 1

        wb.save(PLANILHA_2026)
        wb.close()
        
        app.inserir_log("SUCESSO! ATUALIZAÇÃO CONCLUÍDA.", "VERDE")
        app.inserir_log(f"Total processado: {atualizados} registros.", "BRANCO")
        
        if novos_inseridos:
            app.inserir_log("PROFISSIONAIS ADICIONADOS:", "CIANO")
            for item in novos_inseridos: 
                app.inserir_log(f"   {item}", "VERDE")
        return True 

    except PermissionError:
        app.inserir_log("ERRO: Planilha aberta. Feche e tente novamente.", "VERMELHO")
        return False
    except Exception as e:
        app.inserir_log(f"Erro GERAL: {e}", "VERMELHO")
        return False

# === EVENTOS DO WATCHDOG ===

class RoteadorDownloads(FileSystemEventHandler):
    def __init__(self, app):
        self.app = app
        self.arquivos_em_processamento = set() # Controle de Debounce

    def processar_arquivo(self, caminho_arquivo):
        nome_arquivo = os.path.basename(caminho_arquivo).lower()
        
        # Ignora arquivos temporários (WPS, Excel, Navegador)
        if nome_arquivo.startswith("~") or nome_arquivo.startswith(".") or nome_arquivo.endswith(".crdownload") or nome_arquivo.endswith(".tmp"):
            return

        if nome_arquivo.startswith("emulti") and nome_arquivo.endswith(".pdf"):
            # Evita disparos múltiplos para o mesmo arquivo
            if caminho_arquivo in self.arquivos_em_processamento:
                return
            self.arquivos_em_processamento.add(caminho_arquivo)
            
            self.app.inserir_log(f"ROTEADOR: Arquivo '{nome_arquivo}' detectado em Downloads!", "AZUL")
            if esperar_download_concluir(caminho_arquivo):
                caminho_destino = os.path.join(PASTA_MONITORADA, os.path.basename(caminho_arquivo))
                try:
                    shutil.move(caminho_arquivo, caminho_destino)
                    self.app.inserir_log("ROTEADOR: Arquivo movido para a pasta de produção.", "VERDE")
                except Exception as e:
                    self.app.inserir_log(f"Erro crítico ao mover: {e}", "VERMELHO")
            
            # Limpa o cache após tentar processar
            self.arquivos_em_processamento.discard(caminho_arquivo)

    def on_created(self, event):
        if not event.is_directory: self.processar_arquivo(event.src_path)
            
    def on_moved(self, event):
        if not event.is_directory: self.processar_arquivo(event.dest_path)

class MonitorRelatorios(FileSystemEventHandler):
    def __init__(self, app):
        self.app = app
        self.arquivos_em_processamento = set() # Controle de Debounce

    def processar_arquivo(self, caminho_arquivo):
        nome_arquivo = os.path.basename(caminho_arquivo).lower()
        
        # Ignora arquivos temporários (WPS, Excel, Navegador)
        if nome_arquivo.startswith("~") or nome_arquivo.startswith("."):
            return
            
        if caminho_arquivo.lower().endswith(".pdf"):
            # Evita processar o mesmo arquivo repetidamente de forma simultânea
            if caminho_arquivo in self.arquivos_em_processamento:
                return
            self.arquivos_em_processamento.add(caminho_arquivo)

            self.app.inserir_log("PROCESSADOR: PDF detectado na pasta local. Lendo...", "AMARELO")
            
            if not esperar_download_concluir(caminho_arquivo):
                self.app.inserir_log("ERRO: Tempo limite excedido para o arquivo.", "VERMELHO")
                self.arquivos_em_processamento.discard(caminho_arquivo)
                return

            try:
                tipo, mes, profissionais, duplicados = ler_relatorio_pdf(caminho_arquivo)
                profissionais_planilha = obter_profissionais_planilha()
                
                evento = {
                    "acao": "PROCESSAR_PDF",
                    "caminho": caminho_arquivo,
                    "tipo": tipo,
                    "mes": mes,
                    "profissionais": profissionais,
                    "duplicados": duplicados,
                    "profissionais_planilha": profissionais_planilha
                }
                self.app.fila_eventos.put(evento)

            except Exception as e:
                self.app.inserir_log(f"Erro crítico: {e}", "VERMELHO")
            finally:
                # Permite que o arquivo seja processado de novo no futuro se necessário
                self.arquivos_em_processamento.discard(caminho_arquivo)

    def on_created(self, event):
        if not event.is_directory: self.processar_arquivo(event.src_path)
            
    def on_moved(self, event):
        if not event.is_directory: self.processar_arquivo(event.dest_path)

# === INTERFACE GRÁFICA PRINCIPAL (CUSTOMTKINTER) ===

class AplicacaoGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Sistema de Produção 2026 - Profissional Cátia")
        self.root.geometry("850x650")

        # Configurações globais do CustomTkinter
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.fila_eventos = queue.Queue()

        # Configuração da Caixa de Texto de Log
        self.txt_log = ctk.CTkTextbox(
            self.root, 
            font=("Consolas", 18), 
            state="disabled", 
            wrap="word",
            fg_color="#1e1e1e",
            text_color="#f8f8f2"
        )
        self.txt_log.pack(padx=20, pady=20, fill=tk.BOTH, expand=True)

        self.txt_log.tag_config("VERMELHO", foreground="#ff5555", justify="center")
        self.txt_log.tag_config("VERDE", foreground="#50fa7b", justify="center")
        self.txt_log.tag_config("AMARELO", foreground="#f1fa8c", justify="center")
        self.txt_log.tag_config("AZUL", foreground="#8be9fd", justify="center")
        self.txt_log.tag_config("ROXO", foreground="#bd93f9", justify="center")
        self.txt_log.tag_config("CIANO", foreground="#8be9fd", justify="center")
        self.txt_log.tag_config("BRANCO", foreground="#ffffff", justify="center")
        self.txt_log.tag_config("PADRAO", foreground="#f8f8f2", justify="center")

        self.verificar_fila()
        self.exibir_instrucoes()
        self.iniciar_monitoramento()

    def inserir_log(self, mensagem, cor="PADRAO"):
        self.root.after(0, self._inserir_texto, mensagem, cor)

    def _inserir_texto(self, mensagem, cor):
        self.txt_log.configure(state="normal")
        self.txt_log.insert(tk.END, mensagem + "\n", cor)
        self.txt_log.see(tk.END) 
        self.txt_log.configure(state="disabled")

    def exibir_instrucoes(self):
        self.inserir_log("-" * 70, "ROXO")
        self.inserir_log("SISTEMA DE PRODUÇÃO 2026 - MONITORAMENTO ATIVO", "BRANCO")
        self.inserir_log("-" * 70 + "\n", "ROXO")
        
        self.inserir_log("INSTRUÇÕES AO USUÁRIO:", "AMARELO")
        self.inserir_log("1. MANTENHA ABERTO: Pode minimizar esta janela.", "VERDE")
        self.inserir_log("2. O arquivo precisa começar com 'emulti'.", "VERDE")
        self.inserir_log("3. PLANILHA FECHADA: Não deixe a planilha aberta no Excel/WPS.", "VERMELHO")
        self.inserir_log("4. LIMPEZA: PDFs são apagados após 24h na pasta processados.\n", "VERDE")
        self.inserir_log("Monitorando Downloads e Pasta Local...", "VERDE")

    def verificar_fila(self):
        try:
            while not self.fila_eventos.empty():
                evento = self.fila_eventos.get_nowait()
                if evento["acao"] == "PROCESSAR_PDF":
                    self.solicitar_confirmacao_e_atualizar(evento)
        except queue.Empty:
            pass
        finally:
            self.root.after(200, self.verificar_fila)

    def solicitar_confirmacao_e_atualizar(self, evento):
        tipo = evento["tipo"]
        mes = evento["mes"]
        prof = evento["profissionais"]
        dupl = evento["duplicados"]
        prof_planilha = evento["profissionais_planilha"]
        caminho = evento["caminho"]

        self.inserir_log("\n" + "="*70, "AMARELO")
        self.inserir_log(f" RELATÓRIO DETECTADO: {tipo} | MÊS: {mes} ", "CIANO")
        self.inserir_log("="*70, "AMARELO")
        
        cabecalho = f"{'NOME DO PROFISSIONAL':<40} | {'VALOR':^10} | {'STATUS':^12}"
        self.inserir_log(cabecalho, "BRANCO")
        self.inserir_log("-" * 70, "BRANCO")
        
        novos = False
        for nome, valor in prof.items():
            nome_exibicao = nome[:37] + "..." if len(nome) > 40 else nome
            status = "OK"
            cor = "PADRAO"
            
            if prof_planilha and nome not in prof_planilha:
                status = "NOVO"
                cor = "ROXO"
                novos = True
            elif nome in dupl:
                status = "SOMADO"
                cor = "AZUL"
                
            linha = f"{nome_exibicao:<40} | {valor:^10} | {status:^12}"
            self.inserir_log(linha, cor)
            
        self.inserir_log("="*70 + "\n", "AMARELO")

        msg = f"Relatório {tipo} de {mes} lido com sucesso.\n\nCERTIFIQUE-SE DE QUE A PLANILHA ESTÁ FECHADA ANTES DE CLICAR EM OK!"
        if dupl: msg += "\n\nHá nomes duplicados (valores foram somados)."
        if novos: msg += "\n\nNovos profissionais serão inseridos."

        resposta = messagebox.askokcancel("Ação Necessária", msg, parent=self.root)

        if resposta: 
            self.inserir_log("Iniciando atualização da planilha...", "AMARELO")
            sucesso = atualizar_planilha(tipo, mes, prof, self)
            if sucesso:
                self._limpar_e_arquivar_pdf(caminho)
        else:
            self.inserir_log("Atualização cancelada pelo usuário. O PDF permanecerá na pasta.", "VERMELHO")

    def _limpar_e_arquivar_pdf(self, caminho_arquivo):
        try:
            nome_original = os.path.basename(caminho_arquivo)
            nome_sem_ext, ext = os.path.splitext(nome_original)
            nome_novo = f"{nome_sem_ext}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
            caminho_destino = os.path.join(PASTA_PROCESSADOS, nome_novo)
            
            shutil.move(caminho_arquivo, caminho_destino)
            self.inserir_log("Arquivo processado movido para pasta 'processados'.", "CIANO")
            limpar_arquivos_antigos(PASTA_PROCESSADOS, horas=24, app=self)
            self.inserir_log("Aguardando próximos arquivos...\n", "VERDE")
        except Exception as e:
            self.inserir_log(f"Erro ao mover o PDF: {e}", "VERMELHO")

    def iniciar_monitoramento(self):
        limpar_arquivos_antigos(PASTA_PROCESSADOS, horas=24, app=self)
        
        self.observer = Observer()
        self.observer.schedule(MonitorRelatorios(self), PASTA_MONITORADA, recursive=False)
        self.observer.schedule(RoteadorDownloads(self), PASTA_DOWNLOADS, recursive=False)
        self.observer.start()

    def fechar_aplicacao(self):
        self.observer.stop()
        self.observer.join()
        self.root.destroy()

if __name__ == "__main__":
    root = ctk.CTk()
    app = AplicacaoGUI(root)
    root.protocol("WM_DELETE_WINDOW", app.fechar_aplicacao)
    root.mainloop()


    