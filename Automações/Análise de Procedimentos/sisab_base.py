import os
import re
from openpyxl import load_workbook

from config import Cores
from utils import registrar_log, normalizar_texto
from excel_manager import atualizar_planilha_sisab_memoria, INE_PARA_NOME


class BaseSisabXlsxStrategy:
   
    NOME_LAYOUT = "SISAB"

    def __init__(self, caminho_arquivo, wb_destino, mapa_abas):
        self.caminho_arquivo = caminho_arquivo
        self.wb_destino = wb_destino
        self.mapa_abas = mapa_abas

    def executar(self):
        try:
            registrar_log(
                f"📄 Processando {self.NOME_LAYOUT}: {os.path.basename(self.caminho_arquivo)}",
                Cores.AZUL,
            )

            wb_origem = load_workbook(self.caminho_arquivo, data_only=True)

            for nome_sheet in wb_origem.sheetnames:
                ws = wb_origem[nome_sheet]
                competencia = self._extrair_competencia(ws)
                categoria = self._extrair_categoria(ws)

                registrar_log(f"📅 Competência encontrada: {competencia}", Cores.AZUL)
                registrar_log(f"👨‍⚕️ Categoria encontrada: {categoria}", Cores.VERDE)

                dados_processados = self._extrair_dados_da_aba(ws)

                if not dados_processados:
                    registrar_log(
                        f"⚠️ Nenhum dado compatível encontrado na aba '{nome_sheet}'.",
                        Cores.AMARELO,
                    )
                    continue

                atualizar_planilha_sisab_memoria(
                    dados_processados,
                    competencia,
                    os.path.basename(self.caminho_arquivo),
                    self.wb_destino,
                    self.mapa_abas,
                    profissional_categoria=categoria,
                )

            wb_origem.close()
            registrar_log(f"✅ {self.NOME_LAYOUT} processado com sucesso.", Cores.VERDE)

        except Exception as e:
            registrar_log(f"❌ Erro ao processar {self.NOME_LAYOUT}: {e}", Cores.VERMELHO)

    def _extrair_dados_da_aba(self, ws):
        raise NotImplementedError

    def _localizar_header(self, ws, exige_indicador=False):
        for row in range(1, min(ws.max_row, 80) + 1):
            mapa_colunas = {}
            texto_linha = ""

            for col in range(1, ws.max_column + 1):
                valor = ws.cell(row=row, column=col).value
                if valor is None:
                    continue

                texto = str(valor).strip()
                if not texto:
                    continue

                texto_linha += " | " + texto
                mapa_colunas[texto] = col

            texto_norm = normalizar_texto(texto_linha).lower()
            tem_ine = "equipe ine" in texto_norm or " ine " in f" {texto_norm} "

            if not tem_ine:
                continue

            if exige_indicador:
                tem_indicador = any(
                    termo in texto_norm
                    for termo in (
                        "diabetes",
                        "hipertensao",
                        "pre natal",
                        "puericultura",
                        "puerperio",
                        "saude sexual",
                        "rast cancer",
                        "atendimento domiciliar",
                        "sigtap",
                        "procedimento",
                        "0301100039",
                        "0214010082",
                        "0214010074",
                        "0301100276",
                        "0301100284",
                        "0201020050",
                    )
                )
                if not tem_indicador:
                    continue

            return row, mapa_colunas

        return None, {}

    def _localizar_coluna_ine(self, mapa_colunas):
        for nome_coluna, col in mapa_colunas.items():
            nome_norm = normalizar_texto(nome_coluna).lower()
            if "equipe ine" in nome_norm or nome_norm == "ine" or " ine " in f" {nome_norm} ":
                return col
        return None

    def _extrair_competencia(self, ws):
        for row in range(1, min(ws.max_row, 25) + 1):
            for col in range(1, min(ws.max_column, 15) + 1):
                valor = ws.cell(row=row, column=col).value
                if not valor:
                    continue

                texto = str(valor).strip()
                if "Competência:" in texto:
                    match = re.search(r"Competência:\s*(.+?)(?:\.|$)", texto)
                    if match:
                        return match.group(1).strip()
                    return texto.split(":")[-1].strip()

        return "JANEIRO"

    def _extrair_categoria(self, ws):
        for row in range(1, min(ws.max_row, 25) + 1):
            for col in range(1, min(ws.max_column, 15) + 1):
                valor = ws.cell(row=row, column=col).value
                if not valor:
                    continue

                texto = str(valor).strip()
                if "Categoria Profissional:" in texto:
                    return texto.split(":")[-1].replace(".", "").strip()

        return None

    def _normalizar_ine(self, valor):
        if valor is None:
            return None

        texto = str(valor).strip()
        try:
            numero = str(int(float(texto)))
            return numero[-7:].zfill(7)
        except Exception:
            somente_digitos = re.sub(r"\D", "", texto)
            if not somente_digitos:
                return None
            return somente_digitos[-7:].zfill(7)

    def _converter_numero(self, valor):
        try:
            if valor is None:
                return None

            texto = str(valor).strip()
            if texto == "":
                return None

            texto = texto.replace(".", "").replace(",", ".")
            return float(texto)
        except Exception:
            return None

    def _montar_item(self, ine, indicador, valor):
        nome_posto = INE_PARA_NOME.get(ine, "INE NÃO CADASTRADO")
        return {
            "ine": ine,
            "valor": valor,
            "indicadores": [indicador],
            "nome": nome_posto,
        }
