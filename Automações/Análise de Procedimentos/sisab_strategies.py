from openpyxl import load_workbook

from config import Cores
from utils import registrar_log, normalizar_texto
from sisab_problemas_strategy import ContextoSisabProblemasCondicoes
from sigtap_strategy import ContextoSigtap


class ContextoSisab:
    

    def __init__(self, caminho_arquivo, wb_destino, mapa_abas):
        self.caminho_arquivo = caminho_arquivo
        self.wb_destino = wb_destino
        self.mapa_abas = mapa_abas

    def executar(self):
        tipo_layout = self._identificar_tipo_layout()

        if tipo_layout == "SIGTAP":
            ContextoSigtap(self.caminho_arquivo, self.wb_destino, self.mapa_abas).executar()
            return

        if tipo_layout == "PROBLEMAS_CONDICOES":
            ContextoSisabProblemasCondicoes(self.caminho_arquivo, self.wb_destino, self.mapa_abas).executar()
            return

        registrar_log(
            "⚠️ Tipo de layout SISAB não identificado com segurança. Tentando como Problemas/Condições.",
            Cores.AMARELO,
        )
        ContextoSisabProblemasCondicoes(self.caminho_arquivo, self.wb_destino, self.mapa_abas).executar()

    def _identificar_tipo_layout(self):
        try:
            wb = load_workbook(self.caminho_arquivo, data_only=True, read_only=True)
            textos = []

            for nome_sheet in wb.sheetnames:
                ws = wb[nome_sheet]
                for row in range(1, min(ws.max_row, 20) + 1):
                    for col in range(1, min(ws.max_column, 20) + 1):
                        valor = ws.cell(row=row, column=col).value
                        if valor:
                            textos.append(str(valor))

            wb.close()
            texto_norm = normalizar_texto(" | ".join(textos)).lower()

            if "sigtap" in texto_norm or "tipo de producao procedimento" in texto_norm:
                return "SIGTAP"

            if "probl condicao avaliada" in texto_norm or "tipo de producao atendimento individual" in texto_norm:
                return "PROBLEMAS_CONDICOES"

            # Fallback por códigos SIGTAP no cabeçalho, útil quando a linha do filtro não vem completa.
            if any(codigo in texto_norm for codigo in (
                "0301100039", "0214010082", "0214010074", "0301100276", "0301100284", "0201020050"
            )):
                return "SIGTAP"

            return "DESCONHECIDO"

        except Exception as e:
            registrar_log(f"⚠️ Não foi possível identificar o layout SISAB: {e}", Cores.AMARELO)
            return "DESCONHECIDO"
